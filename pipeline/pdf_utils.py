"""PDF text extraction and cleaning shared by the JD and resume parsers."""
from __future__ import annotations

import re
from pathlib import Path

import fitz  # PyMuPDF

BULLET_CHARS = "•▪◦●○■□➢➤►▶‣⁃∙·-–—*❖✓✔➔→◆◇▸▹»※~"
_NUMBERED_RE = re.compile(r"^\s*(?:\(?\d{1,2}[.)]|\(?[a-zA-Z][.)]|[ivx]{1,4}[.)])\s+")
_PAGE_NUM_RE = re.compile(r"^\s*(page\s*)?\d+(\s*(of|/)\s*\d+)?\s*$", re.I)
_EMAIL_RE = re.compile(r"[\w.+-]+@[\w-]+\.[\w.-]+")
_PHONE_RE = re.compile(r"(\+?\d[\d\s\-()]{7,}\d)")
_URL_RE = re.compile(r"(https?://\S+|www\.\S+|linkedin\.com/\S+|github\.com/\S+)", re.I)


def extract_pdf_text(pdf_path: str | Path) -> tuple[str, float]:
    """
    Return (raw_text, extraction_quality).
    extraction_quality is 1.0 for clean text, 0.5 when the PDF looks
    scanned/garbled (too little text or too many non-ASCII glyphs).
    """
    path = Path(pdf_path)
    if not path.exists():
        raise FileNotFoundError(path)
    suffix = path.suffix.lower()
    if suffix == ".txt":
        text = path.read_text(encoding="utf-8", errors="ignore")
    elif suffix == ".docx":
        text = _extract_docx(path)
    else:
        parts: list[str] = []
        with fitz.open(str(path)) as doc:
            for page in doc:
                parts.append(page.get_text("text"))
        text = "\n".join(parts)

    quality = 1.0
    if len(text.strip()) < 200:
        quality = 0.5
    else:
        non_ascii = sum(1 for ch in text if ord(ch) > 127 and ch not in BULLET_CHARS)
        if non_ascii / max(1, len(text)) > 0.30:
            quality = 0.5
    return text, quality


def _extract_docx(path: Path) -> str:
    """Paragraphs + table cells from a .docx, preserving bullet markers."""
    import docx  # python-docx

    document = docx.Document(str(path))
    lines: list[str] = []
    for para in document.paragraphs:
        txt = para.text.strip()
        if not txt:
            continue
        style = (para.style.name or "").lower() if para.style is not None else ""
        if "list" in style or "bullet" in style:
            txt = "• " + txt
        lines.append(txt)
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                lines.append(" | ".join(dict.fromkeys(cells)))
    return "\n".join(lines)


SUPPORTED_SUFFIXES = (".pdf", ".docx", ".txt")


def clean_text(raw: str) -> str:
    """Normalise PDF artefacts while keeping line structure."""
    text = raw.replace("\r", "\n").replace("\x00", "")
    text = text.replace("\ufb01", "fi").replace("\ufb02", "fl")
    text = re.sub(r"[ \t\u00a0]+", " ", text)
    # de-hyphenate words broken across lines: "develop-\nment"
    text = re.sub(r"(\w)-\n(\w)", r"\1\2", text)
    lines = []
    for line in text.split("\n"):
        line = line.strip()
        if not line or _PAGE_NUM_RE.match(line):
            continue
        lines.append(line)
    # drop lines repeated on many pages (headers/footers)
    counts: dict[str, int] = {}
    for l in lines:
        counts[l] = counts.get(l, 0) + 1
    lines = [l for l in lines if not (counts[l] >= 3 and len(l) < 60)]
    return "\n".join(lines)


def is_contact_line(line: str) -> bool:
    stripped = _EMAIL_RE.sub("", line)
    stripped = _URL_RE.sub("", stripped)
    stripped = _PHONE_RE.sub("", stripped)
    stripped = re.sub(r"[|,;:\s\-•·]+", "", stripped)
    had_contact = bool(_EMAIL_RE.search(line) or _URL_RE.search(line) or _PHONE_RE.search(line))
    return had_contact and len(stripped) <= 3


def strip_contacts(text: str) -> str:
    text = _EMAIL_RE.sub(" ", text)
    text = _URL_RE.sub(" ", text)
    text = _PHONE_RE.sub(" ", text)
    return re.sub(r"[ ]{2,}", " ", text)


def strip_bullet(line: str) -> str:
    s = line.lstrip(BULLET_CHARS + " \t").strip()
    return _NUMBERED_RE.sub("", s, count=1).strip()


def is_bulleted(line: str) -> bool:
    s = line.lstrip(" \t")
    return bool(s) and (s[0] in BULLET_CHARS or bool(_NUMBERED_RE.match(s)))


def looks_like_header(line: str) -> bool:
    """Short, no terminal period, few words, upper-case or colon-terminated."""
    s = line.strip()
    if not s or len(s) > 45:
        return False
    words = s.rstrip(":").split()
    if len(words) > 5:
        return False
    if s.endswith(":"):
        return True
    letters = [c for c in s if c.isalpha()]
    if letters and all(c.isupper() for c in letters):
        return True
    # Title Case short line with no digits / punctuation
    if len(words) <= 3 and all(w[:1].isupper() for w in words) and not re.search(r"[\d,.;()]", s):
        return True
    return False
